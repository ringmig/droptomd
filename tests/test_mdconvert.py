"""Run with the app's Python: ~/Applications/"Drop to MD.app"/Contents/Resources/python/bin/python3 tests/test_mdconvert.py"""
import sys
import tempfile
from pathlib import Path

import pptx
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).parent.parent))
from mdconvert import clean, convert  # noqa: E402

here = Path(__file__)
assert clean("a\x0bb\x0cc", here) == "a\nb\n\nc\n"
assert clean("ﬁelds ‪the‬ ﬂuff", here) == "fields the fluff\n"
assert clean("x\n![Picture 3](Picture3.jpg)\ny", here) == "x\ny\n"
assert clean("![p](https://e.x/p.png)", here) == "![p](https://e.x/p.png)\n"
assert clean(" item", here) == "- item\n"

# A slide with no title placeholder, as Keynote exports it: the largest text becomes the heading
deck = pptx.Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[6])
for text, size, top in [("Raise the\vheartbeat", 60, 1), ("Why we go to work every day", 18, 3)]:
    frame = slide.shapes.add_textbox(Inches(1), Inches(top), Inches(8), Inches(1)).text_frame
    frame.text = text
    frame.paragraphs[0].runs[0].font.size = Pt(size)
with tempfile.TemporaryDirectory() as tmp:
    deck.save(Path(tmp) / "deck.pptx")
    md = convert(Path(tmp) / "deck.pptx")
assert "# Raise the heartbeat\n" in md and "\nWhy we go to work every day" in md, md

assert "Räksmörgås" in convert(here.parent / "fixtures/pdf.pdf")
print("ok")
